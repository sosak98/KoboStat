"""Construction d'un rapport Word (.docx) à partir des éléments d'analyse
sélectionnés par l'utilisateur (tableaux, graphiques, textes d'interprétation)."""
import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_table(doc, df, caption=None):
    if caption:
        p = doc.add_paragraph()
        run = p.add_run(caption)
        run.italic = True
        run.font.size = Pt(10)
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr_cells[i].text = str(col)
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph("")


def build_docx(report_items, title="Résultats et analyses", author=""):
    """report_items: liste de dicts avec clés:
       - type: 'heading' | 'text' | 'table' | 'image'
       - content: str / DataFrame / bytes (PNG)
       - caption (optionnel)
    """
    doc = Document()
    heading = doc.add_heading(title, level=1)
    if author:
        p = doc.add_paragraph(f"Préparé par : {author}")
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for item in report_items:
        t = item["type"]
        if t == "heading":
            doc.add_heading(item["content"], level=item.get("level", 2))
        elif t == "text":
            doc.add_paragraph(item["content"])
        elif t == "table":
            _add_table(doc, item["content"], item.get("caption"))
        elif t == "image":
            doc.add_picture(io.BytesIO(item["content"]), width=Inches(5.5))
            if item.get("caption"):
                p = doc.add_paragraph(item["caption"])
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.italic = True
                    r.font.size = Pt(9)
        doc.add_paragraph("")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
